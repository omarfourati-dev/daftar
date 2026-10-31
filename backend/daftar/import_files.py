"""Reading the uploaded spreadsheet of an import (section 14): ``.xlsx`` (openpyxl), ``.xls`` (xlrd), ``.csv``/``.txt``.

- The format is recognised by content (zip/OLE2/BIFF signatures), not trusted from the file name.
- CSV/TXT: UTF-8 with or without BOM, UTF-16 with BOM (Excel "Texte Unicode"), otherwise Windows-1256 when the
  text is mainly Arabic, otherwise Windows-1252; delimiter ``;`` ``,`` or tab – whichever splits the first lines most
  consistently.
- Limits: 2 MB per file (413), 20 000 data rows and 100 columns per sheet (422). Zip bombs (422
  ``file_too_complex``): every entry of an xlsx at most 5 MB unpacked – ``xl/sharedStrings.xml`` and every worksheet,
  wherever the workbook keeps them (openpyxl loads the shared strings completely, even read-only: security review
  H1) – and 20 MB in total. ``zipfile`` enforces the declared sizes while reading (CRC/size check), so they cannot
  be faked. Cells keep their type (text, int, float, datetime, bool); empty cells are ``None``.
"""

from __future__ import annotations

import codecs
import csv
import io
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from typing import Any

from .import_values import as_text, is_blank

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 20_000
MAX_COLUMNS = 100
MAX_HEADER_ROW = 20
MAX_PART_UNZIPPED = 5 * 1024 * 1024
MAX_UNZIPPED = 20 * 1024 * 1024
CSV_SHEET = "CSV"
_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_BIFF_STREAMS = (b"\x09\x00\x04\x00", b"\x09\x02\x06\x00", b"\x09\x04\x06\x00")
#: Signatures of files that are never text, whatever their name says (PNG, JPEG, GIF, PDF).
_BINARY = (b"\x89PNG", b"\xff\xd8\xff", b"GIF8", b"%PDF")
#: Rows read per sheet: the maximum of data rows plus a possible header block, plus one to detect "too many".
_READ_ROWS = MAX_ROWS + MAX_HEADER_ROW + 1


class ImportFileError(Exception):
    def __init__(self, code: str, message: str, status: int = 422) -> None:
        super().__init__(message)
        self.code, self.message, self.status = code, message, status


@dataclass
class Book:
    kind: str  # xlsx | xls | csv
    sheets: dict[str, list[list[Any]]] = field(default_factory=dict)

    @property
    def sheet_names(self) -> list[str]:
        return list(self.sheets)


@dataclass
class Table:
    headers: list[str]
    rows: list[tuple[int, list[Any]]]  # (row number in the sheet, cells)


def _column_name(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, rest = divmod(index - 1, 26)
        letters = chr(65 + rest) + letters
    return f"Colonne {letters}"


def _trim(row: list[Any] | tuple[Any, ...]) -> list[Any]:
    cells = [None if is_blank(v) else (v.strip() if isinstance(v, str) else v) for v in row]
    while cells and cells[-1] is None:
        cells.pop()
    return cells


# --- formats ---------------------------------------------------------------------------------------------------

def _read_xlsx(data: bytes, max_part: int, max_total: int) -> Book:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            sizes = [info.file_size for info in archive.infolist()]
            if any(size > max_part for size in sizes) or sum(sizes) > max_total:
                raise ImportFileError("file_too_complex", "Fichier trop volumineux une fois décompressé.")
    except zipfile.BadZipFile:
        raise ImportFileError("unreadable_file", "Fichier Excel illisible.") from None
    if "xl/workbook.xml" not in names:
        raise ImportFileError("unsupported_file", "Ce fichier n'est pas un classeur Excel.")
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # openpyxl raises many kinds of errors for broken files
        raise ImportFileError("unreadable_file", "Fichier Excel illisible.") from None
    book = Book("xlsx")
    try:
        for ws in wb.worksheets:
            rows: list[list[Any]] = []
            for row in ws.iter_rows(max_row=_READ_ROWS, max_col=MAX_COLUMNS + 1, values_only=True):
                rows.append(_trim(row))
            book.sheets[ws.title] = rows
    except Exception:
        raise ImportFileError("unreadable_file", "Fichier Excel illisible.") from None
    finally:
        wb.close()
    return book


def _read_xls(data: bytes) -> Book:
    import xlrd

    try:
        wb = xlrd.open_workbook(file_contents=data, logfile=io.StringIO())
    except Exception:
        raise ImportFileError("unreadable_file", "Fichier Excel (.xls) illisible.") from None
    book = Book("xls")
    try:
        for sheet in wb.sheets():
            rows = []
            for r in range(min(sheet.nrows, _READ_ROWS)):
                row: list[Any] = []
                for c in range(min(sheet.ncols, MAX_COLUMNS + 1)):
                    cell = sheet.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            value: Any = xlrd.xldate_as_datetime(cell.value, wb.datemode)
                        except Exception:
                            value = cell.value
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                        value = None
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(cell.value)
                    else:
                        value = cell.value
                    row.append(value)
                rows.append(_trim(row))
            book.sheets[sheet.name] = rows
    finally:
        wb.release_resources()
    return book


def decode_text(data: bytes) -> str:
    if data.startswith(codecs.BOM_UTF8):
        return data[len(codecs.BOM_UTF8):].decode("utf-8", errors="replace")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    arabic = data.decode("cp1256", errors="replace")
    special = [ch for ch in arabic if ord(ch) > 127]
    if special and sum(1 for ch in special if "؀" <= ch <= "ۿ") * 2 >= len(special):
        return arabic
    return data.decode("cp1252", errors="replace")


def detect_delimiter(text: str) -> str:
    lines = [line for line in text.splitlines() if line.strip()][:30]
    best, best_score = ";", (0, 0)
    for delimiter in (";", "\t", ","):
        counts = [len(row) for row in csv.reader(lines, delimiter=delimiter)]
        if not counts or max(counts) < 2:
            continue
        mode, consistent = Counter(counts).most_common(1)[0]
        if mode < 2:
            continue
        score = (consistent, mode)
        if score > best_score:
            best, best_score = delimiter, score
    return best


def _read_csv(data: bytes) -> Book:
    text = decode_text(data).replace("\x00", "")
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=detect_delimiter(text))
    rows: list[list[Any]] = []
    try:
        for row in reader:
            rows.append(_trim(row[: MAX_COLUMNS + 1]))
            if len(rows) >= _READ_ROWS:
                break
    except csv.Error:
        raise ImportFileError("unreadable_file", "Fichier CSV illisible.") from None
    return Book("csv", {CSV_SHEET: rows})


def read_workbook(filename: str, data: bytes, *, max_part: int = MAX_PART_UNZIPPED, max_total: int = MAX_UNZIPPED) -> Book:
    """The sheets of an upload. ``max_part``/``max_total``: unpacked xlsx limits per zip entry and in total (tests
    turn them down instead of building megabyte files)."""
    if len(data) > MAX_BYTES:
        raise ImportFileError("file_too_large", "Fichier trop volumineux (2 Mo au maximum).", status=413)
    if not data.strip():
        raise ImportFileError("empty_file", "Le fichier est vide.")
    if data.startswith(b"PK\x03\x04"):
        book = _read_xlsx(data, max_part, max_total)
    elif data.startswith(_OLE2) or data.startswith(_BIFF_STREAMS):
        book = _read_xls(data)
    elif filename.lower().endswith((".csv", ".txt", ".tsv")) and not data.startswith(_BINARY):
        book = _read_csv(data)
    else:
        raise ImportFileError("unsupported_file", "Format non pris en charge (xlsx, xls, csv ou txt).")
    if not any(any(cell is not None for cell in row) for rows in book.sheets.values() for row in rows):
        raise ImportFileError("empty_file", "Le fichier est vide.")
    return book


# --- one sheet as a table --------------------------------------------------------------------------------------

def read_table(book: Book, sheet: str, header_row: int) -> Table:
    """Headers and the non-empty data rows below ``header_row`` (1-based; 0 = the file has no header row)."""
    if sheet not in book.sheets:
        raise ImportFileError("unknown_sheet", "Feuille introuvable.")
    rows = book.sheets[sheet]
    if not 0 <= header_row <= MAX_HEADER_ROW or header_row > len(rows):
        raise ImportFileError("invalid_header_row", "Ligne d'en-tête invalide.")
    width = max((len(r) for r in rows), default=0)
    if width > MAX_COLUMNS:
        raise ImportFileError("too_many_columns", f"Trop de colonnes ({MAX_COLUMNS} au maximum).")
    head = rows[header_row - 1] if header_row else []
    headers = [as_text(head[i]) if i < len(head) and head[i] is not None else _column_name(i) for i in range(width)]
    data = [
        (number, row + [None] * (width - len(row)))
        for number, row in enumerate(rows[header_row:], start=header_row + 1)
        if any(cell is not None for cell in row)
    ]
    if len(data) > MAX_ROWS:
        raise ImportFileError("too_many_rows", f"Trop de lignes ({MAX_ROWS} au maximum).")
    return Table(headers=headers, rows=data)


def samples(table: Table, index: int, count: int = 3) -> list[str]:
    found = []
    for _, cells in table.rows:
        if cells[index] is not None:
            found.append(as_text(cells[index]))
            if len(found) == count:
                break
    return found
